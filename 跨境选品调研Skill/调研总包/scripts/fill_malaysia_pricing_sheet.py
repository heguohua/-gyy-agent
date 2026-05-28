from copy import copy
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment


ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "data" / "副本TikTok Shop 跨境小店定价表.xlsx"

rows = [
    ["壁挂肥皂盒带沥水", "soap-holder", 2.2, 0.8, 120, 8.9],
    ["牙刷收纳架", "toothbrush-holder", 1.8, 0.6, 90, 7.9],
    ["电动牙刷收纳架", "electric-toothbrush-holder", 4.8, 1.0, 180, 15.9],
    ["牙膏挤压器", "toothpaste-dispenser", 1.5, 0.5, 70, 6.9],
    ["卫生纸架带置物台", "toilet-paper-holder-shelf", 6.8, 1.2, 350, 24.9],
    ["浴室毛巾架", "towel-rack-wall-mounted", 7.2, 1.5, 500, 19.9],
    ["独立式毛巾架", "towel-rack-stand", 15.5, 2.0, 1200, 59.9],
    ["墙面挂钩", "wall-hook", 0.9, 0.4, 40, 5.9],
    ["马桶刷架", "toilet-brush-holder", 5.4, 1.0, 260, 18.9],
    ["鞋架", "shoe-rack", 4.8, 1.2, 650, 12.9],
    ["鞋柜式鞋架", "shoe-rack-cabinet", 12.5, 2.0, 1800, 35.9],
    ["折叠脏衣篮", "foldable-laundry-basket", 5.8, 1.0, 450, 14.9],
    ["带盖脏衣篮", "laundry-basket-with-lid", 9.8, 1.5, 900, 24.9],
    ["衣柜收纳挂袋", "wardrobe-organizer", 4.2, 0.8, 250, 16.9],
    ["折叠衣物收纳盒", "foldable-clothes-storage-box", 3.6, 0.8, 300, 11.9],
    ["床底收纳箱", "under-bed-storage-box", 13.8, 2.2, 1600, 49.9],
    ["包包挂钩", "bag-hook", 0.8, 0.3, 35, 4.9],
    ["围巾挂架", "scarf-hanger-holder", 2.9, 0.6, 110, 9.9],
    ["伸缩晾衣架", "clothes-hanging-rack", 11.8, 1.8, 1300, 29.9],
    ["床边挂篮", "bedside-hanging-basket", 4.1, 0.8, 220, 14.9],
]

history_rows = [
    ["2026-05-27", "免打孔拖把挂架", "mop-holder-wall-mounted", "已推送", "TikTok", "历史去重样本"],
    ["2026-05-27", "缝隙清洁刷", "gap-cleaning-brush", "已推送", "TikTok", "历史去重样本"],
    ["2026-05-27", "浴室免打孔转角架", "bathroom-corner-shelf", "已推送", "TikTok", "历史去重样本"],
    ["2026-05-27", "抽屉分隔收纳盒", "drawer-organizer-box", "已推送", "TikTok", "历史去重样本"],
    ["2026-05-27", "壁挂肥皂盒带沥水", "wall-mounted-soap-holder", "已推送", "TikTok", "历史去重样本"],
    ["2026-05-27", "桌面多格收纳盒", "desktop-storage-box", "已推送", "TikTok", "历史去重样本"],
    ["2026-05-27", "免打孔卷纸架带置物台", "toilet-paper-holder-shelf", "已推送", "TikTok", "历史去重样本"],
    ["2026-05-27", "衣柜挂袋收纳器", "wardrobe-organizer", "已推送", "TikTok", "历史去重样本"],
    ["2026-05-27", "床底收纳袋", "under-bed-storage-bag", "已推送", "TikTok", "历史去重样本"],
    ["2026-05-27", "滚筒粘毛器", "lint-roller", "已推送", "TikTok", "历史去重样本"],
]


def copy_style(src, dst):
    dst._style = copy(src._style)
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = copy(src.number_format)
        dst.protection = copy(src.protection)


wb = load_workbook(PATH)
ws = wb["马来西亚"]

if ws.max_column < 14:
    ws.insert_cols(13, 2)

ws["M2"] = "推送优先级"
ws["N2"] = "是否重复"
ws["M3"] = "排序"
ws["N3"] = "历史去重"

priority_order = {
    "wall-hook": 1,
    "toothbrush-holder": 2,
    "toothpaste-dispenser": 3,
    "soap-holder": 4,
    "wall-mounted-soap-holder": 4,
    "wardrobe-organizer": 5,
    "foldable-clothes-storage-box": 6,
    "bedside-hanging-basket": 7,
    "bag-hook": 8,
    "scarf-hanger-holder": 9,
    "toilet-brush-holder": 10,
    "electric-toothbrush-holder": 11,
    "foldable-laundry-basket": 12,
    "towel-rack-wall-mounted": 13,
    "shoe-rack": 14,
    "toilet-paper-holder-shelf": 15,
    "clothes-hanging-rack": 16,
    "laundry-basket-with-lid": 17,
    "shoe-rack-cabinet": 18,
    "under-bed-storage-box": 19,
    "towel-rack-stand": 20,
}

# Preserve the existing top layout, only fill the data body.
for r in range(5, 25):
    for c in range(1, 15):
        ws.cell(r, c).value = None

for idx, row in enumerate(rows, start=5):
    name, sku, cost, pack, weight, display_price = row
    ws[f"A{idx}"] = name
    ws[f"B{idx}"] = sku
    ws[f"C{idx}"] = cost
    ws[f"D{idx}"] = pack
    ws[f"E{idx}"] = weight
    ws[f"F{idx}"] = display_price
    ws[f"G{idx}"] = f"=K{idx}/F{idx}"
    ws[f"H{idx}"] = f"=(C{idx}+D{idx})/1.5"
    ws[f"I{idx}"] = f"=E{idx}/10*0.15"
    ws[f"J{idx}"] = f"=F{idx}*0.1"
    ws[f"K{idx}"] = f"=F{idx}-J{idx}-H{idx}-I{idx}"
    ws[f"L{idx}"] = f"=K{idx}*1.5"
    ws[f"M{idx}"] = priority_order.get(sku, 99)
    ws[f"N{idx}"] = f'=IF(COUNTIF(\'每日去重表\'!$C:$C,B{idx})>0,"重复","不重复")'

if "每日去重表" in wb.sheetnames:
    del wb["每日去重表"]
history_ws = wb.create_sheet("每日去重表")
history_headers = ["日期", "产品名", "sku", "是否已推送", "推送渠道", "备注"]
history_ws.append(history_headers)
for row in history_rows:
    history_ws.append(row)

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for wsx in [ws, history_ws]:
    for cell in wsx[2]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in wsx[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Extend basic styles from row 5 to new rows so the sheet stays consistent.
for r in range(5, 25):
    for c in range(1, 15):
        copy_style(ws.cell(5, c), ws.cell(r, c))

for r in range(2, history_ws.max_row + 1):
    for c in range(1, history_ws.max_column + 1):
        history_ws.cell(r, c).border = border
        history_ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
        if r == 1:
            history_ws.cell(r, c).fill = header_fill
            history_ws.cell(r, c).font = header_font

for col, width in {
    "A": 26, "B": 14, "C": 12, "D": 12, "E": 12, "F": 14,
    "G": 10, "H": 12, "I": 12, "J": 12, "K": 12, "L": 12,
    "M": 12, "N": 12
}.items():
    ws.column_dimensions[col].width = width

for col, width in {"A": 14, "B": 18, "C": 24, "D": 12, "E": 12, "F": 24}.items():
    history_ws.column_dimensions[col].width = width

ws.freeze_panes = "A5"
history_ws.freeze_panes = "A2"
ws.auto_filter.ref = "A2:N24"
history_ws.auto_filter.ref = f"A1:F{history_ws.max_row}"

wb.save(PATH)
print(PATH)
