from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path


products = [
    {
        "产品名": "免打孔拖把挂架",
        "英文名": "Wall Mounted Mop Holder",
        "类目": "清洁工具收纳",
        "是否厨房": "否",
        "来源平台": "Lazada Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "0.68-15.90",
        "建议售价(RM)": "6.90-12.90",
        "趋势判断": "稳定需求型",
        "内容角度": "收纳前后对比、浴室/阳台整洁",
        "无货源适配度": "高",
        "风险等级": "中低",
        "推荐等级": "优先测试",
        "去重标识": "mop-holder-wall-mounted",
        "建议标题": "Wall Mounted Mop Holder No Drill Storage Rack for Bathroom",
        "商家/商品链接": "https://www.lazada.com.my/videodetail/?video_id=30045880169",
        "备注": "适合做前后对比内容；承重和粘性要重点筛。",
    },
    {
        "产品名": "缝隙清洁刷",
        "英文名": "Gap Cleaning Brush",
        "类目": "清洁工具",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "1.53-3.29",
        "建议售价(RM)": "4.90-8.90",
        "趋势判断": "稳定需求型",
        "内容角度": "浴室角落、窗槽、地缝清洁前后对比",
        "无货源适配度": "高",
        "风险等级": "中低",
        "推荐等级": "优先测试",
        "去重标识": "gap-cleaning-brush",
        "建议标题": "Gap Cleaning Brush for Bathroom Corner Deep Cleaning",
        "商家/商品链接": "https://shopee.com.my/Deep-Clean-Scrub-Brushes-Window-Toilet-Ceramic-Tile-Gap-Cleaning-Brush-Scraper-With-Long-Handle-Corner-Groove-Crevice-Cleaning-Brush-For-Bathroom-And-Kitchen-Surface-Household-mol-i.252747406.23781140666",
        "备注": "问题解决型很强，适合短视频污渍清理演示。",
    },
    {
        "产品名": "浴室免打孔转角架",
        "英文名": "Bathroom Corner Shelf",
        "类目": "浴室收纳",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "8.00-24.00",
        "建议售价(RM)": "12.90-22.90",
        "趋势判断": "稳定需求型",
        "内容角度": "浴室整洁、瓶瓶罐罐收纳前后对比",
        "无货源适配度": "高",
        "风险等级": "中低",
        "推荐等级": "优先测试",
        "去重标识": "bathroom-corner-shelf",
        "建议标题": "Bathroom Corner Shelf No Drill Storage Rack for Shower",
        "商家/商品链接": "https://shopee.com.my/list/bathroom/corner/shelf",
        "备注": "浴室收纳是非厨房居家日用里较稳的内容方向。",
    },
    {
        "产品名": "抽屉分隔收纳盒",
        "英文名": "Drawer Organizer",
        "类目": "家居收纳",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "5.00-19.00",
        "建议售价(RM)": "9.90-16.90",
        "趋势判断": "稳定需求型",
        "内容角度": "抽屉整理、桌面分类收纳",
        "无货源适配度": "高",
        "风险等级": "中低",
        "推荐等级": "优先测试",
        "去重标识": "drawer-organizer-box",
        "建议标题": "Drawer Organizer Box for Home Storage and Easy Sorting",
        "商家/商品链接": "https://shopee.com.my/list/drawer/organizer",
        "备注": "标准化高，适合做多格/多尺寸组合。",
    },
    {
        "产品名": "多层收纳推车",
        "英文名": "Slim Storage Trolley",
        "类目": "家居收纳",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "25.00-79.00",
        "建议售价(RM)": "39.90-69.90",
        "趋势判断": "稳定需求型",
        "内容角度": "浴室/卧室小空间收纳前后对比",
        "无货源适配度": "中低",
        "风险等级": "中",
        "推荐等级": "观察后可测",
        "去重标识": "slim-storage-trolley",
        "建议标题": "Slim Storage Trolley for Bathroom Bedroom Small Space Organizer",
        "商家/商品链接": "https://shopee.com.my/list/storage/trolley",
        "备注": "内容直观但体积较大，更适合后续测。",
    },
    {
        "产品名": "壁挂肥皂盒带沥水",
        "英文名": "Wall Mounted Soap Holder",
        "类目": "浴室收纳",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "3.00-10.00",
        "建议售价(RM)": "6.90-9.90",
        "趋势判断": "稳定需求型",
        "内容角度": "肥皂不积水、浴室整洁",
        "无货源适配度": "高",
        "风险等级": "低",
        "推荐等级": "优先测试",
        "去重标识": "wall-mounted-soap-holder",
        "建议标题": "Wall Mounted Soap Holder with Drain Tray for Bathroom",
        "商家/商品链接": "https://shopee.com.my/list/soap/holder",
        "备注": "低售后标准品，适合搭配浴室收纳系列。",
    },
    {
        "产品名": "桌面多格收纳盒",
        "英文名": "Desktop Storage Box",
        "类目": "家居收纳",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "6.00-18.00",
        "建议售价(RM)": "9.90-15.90",
        "趋势判断": "稳定需求型",
        "内容角度": "桌面清爽、文具化妆品分区收纳",
        "无货源适配度": "高",
        "风险等级": "中低",
        "推荐等级": "优先测试",
        "去重标识": "desktop-storage-box",
        "建议标题": "Desktop Storage Box for Home Office Daily Use Organizer",
        "商家/商品链接": "https://shopee.com.my/list/desktop/storage/box",
        "备注": "泛用性强，适合女生桌面/租房整理场景。",
    },
    {
        "产品名": "滚筒粘毛器",
        "英文名": "Lint Roller",
        "类目": "清洁工具",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "2.00-9.00",
        "建议售价(RM)": "4.90-8.90",
        "趋势判断": "稳定需求型",
        "内容角度": "衣物/床品/沙发去毛演示",
        "无货源适配度": "高",
        "风险等级": "低",
        "推荐等级": "观察后可测",
        "去重标识": "lint-roller",
        "建议标题": "Lint Roller for Clothes Sofa and Pet Hair Cleaning",
        "商家/商品链接": "https://shopee.com.my/list/lint/roller",
        "备注": "需求稳定，但爆点弱于强对比型收纳/清洁产品。",
    },
    {
        "产品名": "免打孔卷纸架带置物台",
        "英文名": "Toilet Paper Holder with Shelf",
        "类目": "浴室收纳",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "8.00-22.00",
        "建议售价(RM)": "12.90-19.90",
        "趋势判断": "稳定需求型",
        "内容角度": "浴室更整洁、手机可放置",
        "无货源适配度": "高",
        "风险等级": "中低",
        "推荐等级": "优先测试",
        "去重标识": "toilet-paper-holder-shelf",
        "建议标题": "Toilet Paper Holder with Shelf No Drill Bathroom Organizer",
        "商家/商品链接": "https://shopee.com.my/list/toilet/paper/holder/shelf",
        "备注": "浴室场景明确，适合免打孔收纳方向。",
    },
    {
        "产品名": "衣柜挂袋收纳器",
        "英文名": "Hanging Closet Organizer",
        "类目": "衣柜收纳",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "10.00-28.00",
        "建议售价(RM)": "15.90-24.90",
        "趋势判断": "稳定需求型",
        "内容角度": "衣柜分层整理、小空间利用",
        "无货源适配度": "高",
        "风险等级": "中低",
        "推荐等级": "优先测试",
        "去重标识": "hanging-closet-organizer",
        "建议标题": "Hanging Closet Organizer for Wardrobe Clothes Storage",
        "商家/商品链接": "https://shopee.com.my/list/hanging/closet/organizer",
        "备注": "适合租房、小空间、衣柜整理方向。",
    },
    {
        "产品名": "折叠脏衣篮",
        "英文名": "Foldable Laundry Basket",
        "类目": "家居收纳",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "12.00-35.00",
        "建议售价(RM)": "19.90-29.90",
        "趋势判断": "稳定需求型",
        "内容角度": "脏衣整理、小户型收纳",
        "无货源适配度": "中",
        "风险等级": "中",
        "推荐等级": "观察后可测",
        "去重标识": "foldable-laundry-basket",
        "建议标题": "Foldable Laundry Basket for Bedroom Bathroom Storage",
        "商家/商品链接": "https://shopee.com.my/list/foldable/laundry/basket",
        "备注": "需求清晰但体积偏大，适合后续优化版本。",
    },
    {
        "产品名": "床底收纳袋",
        "英文名": "Under Bed Storage Bag",
        "类目": "卧室收纳",
        "是否厨房": "否",
        "来源平台": "Shopee Malaysia",
        "市场": "马来西亚",
        "对标价格(RM)": "9.00-26.00",
        "建议售价(RM)": "14.90-22.90",
        "趋势判断": "稳定需求型",
        "内容角度": "隐藏收纳、换季衣物整理",
        "无货源适配度": "高",
        "风险等级": "中低",
        "推荐等级": "优先测试",
        "去重标识": "under-bed-storage-bag",
        "建议标题": "Under Bed Storage Bag for Clothes Blanket Space Saving",
        "商家/商品链接": "https://shopee.com.my/list/under/bed/storage",
        "备注": "卧室收纳需求稳定，适合空间优化内容。",
    },
]

ROOT = Path(__file__).resolve().parents[1]
output_path = ROOT / "outputs" / "马来西亚非厨房居家日用-优先测试与每日去重版.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "优先测试池"

summary = wb.create_sheet("使用说明")
daily = wb.create_sheet("每日推送去重表")

summary["A1"] = "文件说明"
summary["A1"].font = Font(bold=True, size=14)
summary["A3"] = "筛选条件"
summary["B3"] = "仅保留马来西亚居家日用；不包含厨房类产品；以优先测试为主。"
summary["A4"] = "去重规则"
summary["B4"] = "同一个去重标识的产品，推送记录表中若已标记“已推送”，后续不再重复推送。"
summary["A5"] = "建议操作"
summary["B5"] = "每天先从优先测试池筛选未推送产品，再登记到每日推送去重表。"
summary.column_dimensions["A"].width = 16
summary.column_dimensions["B"].width = 90

headers = list(products[0].keys())
ws.append(headers)
for item in products:
    ws.append([item[h] for h in headers])

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.border = border
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)

widths = {
    "A": 18, "B": 28, "C": 16, "D": 10, "E": 18, "F": 12, "G": 14, "H": 14,
    "I": 14, "J": 28, "K": 12, "L": 10, "M": 12, "N": 24, "O": 52, "P": 70, "Q": 36
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

for row_idx in range(2, ws.max_row + 1):
    grade = ws[f"M{row_idx}"].value
    if grade == "优先测试":
        fill = PatternFill("solid", fgColor="E2F0D9")
    else:
        fill = PatternFill("solid", fgColor="FFF2CC")
    for col_idx in range(1, ws.max_column + 1):
        ws.cell(row=row_idx, column=col_idx).fill = fill

daily_headers = ["日期", "产品名", "去重标识", "是否已推送", "推送渠道", "内容方向", "备注"]
daily.append(daily_headers)
for cell in daily[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.border = border
    cell.alignment = Alignment(horizontal="center", vertical="center")

seed_rows = [
    ["2026-05-28", "免打孔拖把挂架", "mop-holder-wall-mounted", "已推送", "TikTok", "浴室/阳台收纳前后对比", "示例记录，后续同去重标识不重复推送"],
    ["2026-05-29", "缝隙清洁刷", "gap-cleaning-brush", "待推送", "TikTok", "角落污渍清洁前后对比", "发布后把状态改成已推送"],
]
for r in seed_rows:
    daily.append(r)

for row in daily.iter_rows(min_row=2):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)

for col, width in {"A": 14, "B": 18, "C": 28, "D": 12, "E": 12, "F": 30, "G": 40}.items():
    daily.column_dimensions[col].width = width
daily.freeze_panes = "A2"
daily.auto_filter.ref = f"A1:G{daily.max_row}"

wb.save(output_path)
print(output_path)
